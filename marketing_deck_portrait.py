"""行銷簡報「A4 直式」：直式照片多的物件(透天、店面)用，印成 A4 直式 DM。
保留富住通元素：頁首/頁尾公版外框(含 LOGO、Full Life 標語)、封面 LOGO＋紅色售/租章、封底直式 LOGO＋Full Life＋公司電話。
頁面依直式重新排版(不是把橫式壓扁)：封面上 2/3 滿版照片、物件資料、樓層配置、地理位置、地籍圖、成果圖、格局圖、
照片(直式兩張並排、橫式一張滿寬)、市場行情、封底。
共用 marketing_deck 的畫圖工具(M＝marketing_deck 模組)；進來時把 M 的頁面座標改成直式實際尺寸(SX=SY=1)。"""
import os
import re
import glob


def build(M, case_dir, base_dir, d, exp, chosen, out_path, f, show_agents=True, back_cover=True, sel=None, template="building"):
    from pptx.util import Inches
    from PIL import Image
    # ---- 直式頁面座標：設計單位＝實際吋 ----
    M.SW, M.SH = M.REAL_W, M.REAL_H
    M.SX = M.SY = 1.0
    M.FOOTER_Y = M.SH - Inches(0.27)
    W, H = M.SW, M.SH
    L, R = Inches(0.45), W - Inches(0.45)
    CW = R - L
    TOP = Inches(1.15)                    # 頁首(公版外框 1 吋)下面
    BOT = M.FOOTER_Y - Inches(0.12)
    tmp = M.tempfile.mkdtemp(prefix="deck_")
    deck = M.Deck(f, tmp)                 # 會依直式尺寸重組富住通外框

    def new(title=None):
        s = deck.prs.slides.add_slide(deck.blank)
        deck.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = M.BG
        if title:
            if deck.frame:
                s.shapes.add_picture(deck.frame, 0, 0, M.REAL_W, M.REAL_H)
            else:
                M.box(s, 0, 0, W, Inches(1.0), M.SKY, radius=False)
            M.text(s, Inches(0.4), Inches(0.12), Inches(4.3), Inches(0.76), title, 26, True, M.WHITE, anchor=M.MSO_ANCHOR.MIDDLE)
            M.text(s, W - Inches(1.3), M.FOOTER_Y, Inches(1.0), H - M.FOOTER_Y, str(deck.n), 10, True, M.WHITE,
                   M.PP_ALIGN.RIGHT, M.MSO_ANCHOR.MIDDLE)
        return s
    deck.new = new

    roots = [r for r in (case_dir, exp) if r]
    by = {c: M.ordered([p for p in chosen if p["cat"] == c]) for c in M.CATS}
    cover = (by["封面"] or by[M.EXT_CAT] or by[M.STREET_CAT] or by["陽台景觀"] or by["客廳"] or [None])[0]
    rent = f["deal"] == "租"

    # ---- 資料來源(同橫式：使用者在「頁面內容」選的優先) ----
    def first(pattern):
        for r in roots:
            hits = sorted(glob.glob(os.path.join(r, pattern)))
            if hits:
                return hits[0]
        return None
    plans = []
    for r in roots:
        hits = sorted(glob.glob(os.path.join(r, "0.謄本", "*建物測量成果圖*.pdf")))
        if hits:
            plans = hits[:1]
            break
    if not plans:
        for r in roots:
            hits = sorted(h for h in glob.glob(os.path.join(r, "0.謄本", "*建物測量成果圖*"))
                          if h.lower().endswith((".png", ".jpg", ".jpeg")))
            if hits:
                plans = hits[:3]
                break
    plans = M.sel_paths(sel, "plan", plans)
    cad = first(os.path.join("0.謄本", "*地籍圖*標註*.pdf")) or first(os.path.join("0.謄本", "*地籍圖*.pdf"))
    cad_pages = M.sel_paths(sel, "cadastral", [cad] if cad else [])
    layouts = [p["path"] for p in by[M.LAYOUT_CAT]]
    names = {os.path.basename(x).lower() for x in layouts}
    layouts += [x for x in M.sel_paths(sel, "layout", M.find_layouts(roots)) if os.path.basename(x).lower() not in names]

    # 面寬/深度：物件編輯器 → 成果圖 → 地籍圖
    f["dims_show"] = f["dims"]
    if not f["dims"] and plans and plans[0].lower().endswith(".pdf"):
        try:
            big = max([v for _, v in f.get("floors") or []] or [M._num(f.get("main")) or 0])
            pd_ = M.plan_dims(plans[0], big * 3.30579 if big else None)
            if pd_:
                f["dims_show"] = f"約 {pd_[0]:g}／{pd_[1]:g} 米（成果圖量測）"
        except Exception:
            pass
    if not f["dims_show"] and cad_pages:
        try:
            dm = M.parcel_dims(cad_pages[0])
            if dm:
                f["dims_show"] = f"約 {dm['frontage']:g}／{dm['depth']:g} 米（地籍圖量測）"
        except Exception:
            pass
    phones = M.lookup_phones(f["agents"], base_dir)

    # ================================================================ 封面
    M.log("製作：封面（A4 直式）")
    s = new()
    ph = Inches(7.3)
    img = None
    if cover:
        _, img = M.picture(s, cover["path"], 0, 0, W, ph, tmp)
    else:
        M.box(s, 0, 0, W, ph, M.SKY_DK, radius=False)
    st = M.box(s, Inches(0.4), Inches(0.4), Inches(1.05), Inches(1.05), M.STAMP_RED)          # 紅色售/租章
    st.adjustments[0] = 0.12
    M.text(s, Inches(0.4), Inches(0.4), Inches(1.05), Inches(1.05), f["deal"], 54, True, M.WHITE, M.PP_ALIGN.CENTER,
           M.MSO_ANCHOR.MIDDLE)
    logo = M.asset(M.ASSET_LOGO)                                                                # 右上富住通 LOGO
    if logo:
        lw = Inches(3.3)
        iw, ih = Image.open(logo).size
        lh = int(lw * ih / iw)
        M.box(s, W - lw - Inches(0.6), Inches(0.38), lw + Inches(0.3), lh + Inches(0.24),
              M.tone(img, (0.6, 0.0, 1.0, 0.25)) if img else M.SKY_DK, alpha=0.55)
        s.shapes.add_picture(logo, W - lw - Inches(0.45), Inches(0.5), lw, lh)
    # 下方資訊區
    M.box(s, 0, ph, W, H - ph, M.WHITE, radius=False)
    M.box(s, 0, ph, W, Inches(0.09), M.SKY, radius=False)
    y = ph + Inches(0.3)
    M.text(s, L, y, CW, Inches(0.95), f["title"], 26 if len(f["title"]) <= 16 else 21, True, M.INK)
    M.text(s, L, y + Inches(0.9), CW, Inches(0.4), f["door"], 16, True, M.SKY_DK)
    M.gold(s, L, y + Inches(1.4), Inches(4.0), Inches(0.9))
    M.text(s, L, y + Inches(1.4), Inches(4.0), Inches(0.9),
           f"{'月租' if rent else '開價'} {f['price']} {f['price_unit']}", 28, True, M.WHITE, M.PP_ALIGN.CENTER,
           M.MSO_ANCHOR.MIDDLE)
    stats = [(k, v) for k, v in (("權狀坪數", f"{f['area']} 坪" if str(f.get("area", "")).strip() else ""),
                                 ("樓層", f.get("floor")), ("面寬／深度", _short_dims(f.get("dims_show"))),
                                 ("面臨道路", f.get("road_w")), ("格局", f.get("layout")), ("屋齡", f.get("age"))) if v][:4]
    sx = L + Inches(4.25)
    sw_ = R - sx
    for i, (k, v) in enumerate(stats):
        cx = sx + (sw_ // 2) * (i % 2)
        cy = y + Inches(1.3) + Inches(0.58) * (i // 2)
        M.text(s, cx, cy, sw_ // 2, Inches(0.25), k, 11, True, M.SKY_DK)
        M.text(s, cx, cy + Inches(0.24), sw_ // 2, Inches(0.32), v, 15, True, M.INK)
    agents = [f"{nm} {phones.get(nm, '')}".strip() for nm in f["agents"][:3]] if show_agents else []
    if agents:
        M.box(s, L, H - Inches(0.95), CW, Inches(0.015), M.LINE, radius=False)
        M.text(s, L, H - Inches(0.88), CW, Inches(0.5), "專案人員：" + "　".join(agents), 15, True, M.INK,
               anchor=M.MSO_ANCHOR.MIDDLE)
    M.text(s, L, H - Inches(0.42), CW, Inches(0.3), f"{M.COMPANY}　{M.COMPANY_TEL}", 11, True, M.SKY_DK,
           anchor=M.MSO_ANCHOR.MIDDLE)

    # ================================================================ 物件資料
    if M.sel_on(sel, "overview"):
        M.log("製作：物件資料")
        s = new("物件資料")
        alt = [(k, v) for k, v in (("格局", f.get("layout")), ("面臨道路", f.get("road_w")), ("土地坪數",
               f"{f['land']} 坪" if str(f.get("land", "")).strip() else ""), ("使用分區", f.get("zone"))) if str(v or "").strip()]
        cards = [(f["price_label"], f["price"], f["price_unit"]), ("權狀坪數", f["area"], "坪"),
                 ("樓層", f.get("floor") or "—", ""), (alt[0][0], alt[0][1], "") if alt else ("格局", "—", ""),
                 ("面寬／深度", _short_dims(f.get("dims_show")) or "—", ""), ("屋齡", f.get("age") or "—", "")]
        cw, chh, g = (CW - Inches(0.2)) // 2, Inches(1.05), Inches(0.14)
        for i, (k, v, u) in enumerate(cards):
            x, yy = L + (cw + Inches(0.2)) * (i % 2), TOP + (chh + g) * (i // 2)
            M.box(s, x, yy, cw, chh, M.WHITE, M.LINE)
            if i == 0:
                M.gold(s, x, yy, Inches(0.1), chh, radius=False)
            else:
                M.box(s, x, yy, Inches(0.1), chh, M.SKY, radius=False)
            M.text(s, x + Inches(0.25), yy + Inches(0.07), cw - Inches(0.3), Inches(0.32), k, 14, True, M.INK)
            tb = M.text(s, x + Inches(0.25), yy + Inches(0.4), cw - Inches(0.3), Inches(0.58), str(v),
                        28 if len(str(v)) <= 7 else 19, True, M.GOLD_LO if i == 0 else M.NUM, anchor=M.MSO_ANCHOR.MIDDLE)
            if u:
                r_ = tb.text_frame.paragraphs[0].add_run()
                r_.text = f" {u}"
                M._font(r_, 13, True, M.INK)
        yy = TOP + (chh + g) * 3 + Inches(0.05)
        rows = [("主建物", f"{f['main']} 坪"), ("附屬建物", f"{f['attach']} 坪"), ("公共設施", f"{f['public']} 坪"),
                ("土地坪數", f"{f['land']} 坪"), ("面臨道路", f.get("road_w")), ("使用分區", f.get("zone")),
                ("建蔽／容積", f.get("bcr_far")), ("建築結構", f.get("struct")), ("座向", f.get("facing")),
                ("適合行業", "、".join(f.get("trades") or [])), ("鄰近學校", f.get("school")), ("所在商圈", f.get("district"))]
        rows = [(k, v) for k, v in rows if str(v or "").strip() and not str(v).strip().startswith(" 坪")
                and str(v).strip() != "坪" and not M._zero(v)][:10]
        rows = [(k, v) for k, v in rows if k != (alt[0][0] if alt else "")]
        half = (len(rows) + 1) // 2
        pts = f["points"][:5] or ["（請在物件編輯器填寫「訴求重點」）"]
        # 下面空間依內容分配：明細表每列最高 0.6 吋、特色每點最高 0.95 吋(不留一大片空白)
        room = BOT - yy - Inches(0.75)
        rh = min(Inches(0.6), max(Inches(0.42), int(room * 0.45) // max(1, half)))
        M.box(s, L, yy, CW, rh * half + Inches(0.2), M.PANEL)
        for i, (k, v) in enumerate(rows):
            cx = L + Inches(0.2) + (CW // 2) * (i // half)
            cy = yy + Inches(0.1) + rh * (i % half)
            M.text(s, cx, cy, Inches(1.25), rh, k, 14, True, M.SKY_DK, anchor=M.MSO_ANCHOR.MIDDLE)
            M.text(s, cx + Inches(1.25), cy, CW // 2 - Inches(1.5), rh, str(v), 15 if len(str(v)) <= 11 else 12, True,
                   M.INK, anchor=M.MSO_ANCHOR.MIDDLE)
        yy += rh * half + Inches(0.4)
        M.text(s, L, yy, CW, Inches(0.42), "物件特色", 18, True, M.SKY_DK)
        yy += Inches(0.5)
        each = min(Inches(0.95), (BOT - yy) // max(1, len(pts)))
        for i, p_ in enumerate(pts):
            M.gold(s, L, yy + (each - Inches(0.44)) // 2, Inches(0.44), Inches(0.44), oval=True)
            M.text(s, L, yy + (each - Inches(0.44)) // 2, Inches(0.44), Inches(0.44), str(i + 1), 15, True, M.WHITE,
                   M.PP_ALIGN.CENTER, M.MSO_ANCHOR.MIDDLE)
            M.box(s, L + Inches(0.6), yy + Inches(0.05), CW - Inches(0.6), each - Inches(0.1), M.WHITE, M.LINE)
            M.text(s, L + Inches(0.75), yy, CW - Inches(0.9), each, p_, 16 if len(p_) <= 30 else 13, True, M.INK,
                   anchor=M.MSO_ANCHOR.MIDDLE)
            yy += each

    # ================================================================ 樓層配置(透天)
    if f.get("floors") and len(f["floors"]) > 1 and M.sel_on(sel, "floors"):
        M.log("製作：樓層配置")
        s = new("樓層配置")
        fl = sorted(f["floors"], key=lambda x: M._floor_order(x[0]), reverse=True)
        vmax = max(v for _, v in fl)
        rh = min(Inches(1.35), Inches(6.2) // len(fl))
        for i, (k, v) in enumerate(fl):
            yy = TOP + Inches(0.1) + rh * i
            M.box(s, L, yy + Inches(0.08), Inches(1.6), rh - Inches(0.16), M.PANEL, M.SKY)
            M.text(s, L, yy + Inches(0.08), Inches(1.6), rh - Inches(0.16), k, 20, True, M.SKY_DK, M.PP_ALIGN.CENTER,
                   M.MSO_ANCHOR.MIDDLE)
            bw = int((CW - Inches(3.4)) * v / vmax)
            M.box(s, L + Inches(1.75), yy + Inches(0.2), max(bw, Inches(0.2)), rh - Inches(0.4), M.SKY)
            M.text(s, L + Inches(1.9) + bw, yy, Inches(1.5), rh, f"{v:g} 坪", 20, True, M.INK, anchor=M.MSO_ANCHOR.MIDDLE)
        yy = TOP + Inches(0.2) + rh * len(fl)
        M.gold(s, L, yy, CW, Inches(0.75))
        M.text(s, L + Inches(0.3), yy, CW, Inches(0.75), f"權狀合計　{f['area']} 坪", 22, True, M.WHITE,
               anchor=M.MSO_ANCHOR.MIDDLE)
        info = [(k, v) for k, v in (("總樓層", f"地上 {f['above']} 層" + (f"・地下 {f['below']} 層" if f.get("below") else "")
                                     if f.get("above") else ""), ("建築結構", f.get("struct")),
                                    ("竣工", re.sub(r"(?<=[國年月])0+(?=\d)", "", str(f.get("built") or ""))),
                                    ("格局", f.get("layout"))) if str(v or "").strip()]
        yy += Inches(0.95)
        info = info + [(k, v) for k, v in (("面寬／深度", _short_dims(f.get("dims_show"))), ("面臨道路", f.get("road_w")),
                                           ("土地坪數", f"{f['land']} 坪" if str(f.get("land", "")).strip() else ""))
                       if str(v or "").strip()]
        ih = min(Inches(0.7), (BOT - yy) // max(1, len(info)))
        for k, v in info:
            M.box(s, L, yy, CW, ih - Inches(0.08), M.WHITE, M.LINE)
            M.text(s, L + Inches(0.25), yy, Inches(1.6), ih - Inches(0.08), k, 15, True, M.SKY_DK, anchor=M.MSO_ANCHOR.MIDDLE)
            M.text(s, L + Inches(1.9), yy, CW - Inches(2.1), ih - Inches(0.08), str(v), 17, True, M.INK,
                   anchor=M.MSO_ANCHOR.MIDDLE)
            yy += ih

    # ================================================================ 地理位置
    ll = M.case_coords(case_dir, base_dir, d) if M.sel_on(sel, "location") else None
    if ll:
        M.log("製作：地理位置")
        mh = Inches(5.6)
        big = M.tile_map(ll[0], ll[1], 16, int(CW / 914400 * 200), int(mh / 914400 * 200), "EMAP",
                         rings=((500, "500 公尺・步行約 6 分"), (1000, "1 公里")))
        if big is not None:
            s = new("地理位置")
            pb = os.path.join(tmp, "loc_map.png")
            big.save(pb, optimize=True)
            M.box(s, L - Inches(0.05), TOP - Inches(0.05), CW + Inches(0.1), mh + Inches(0.1), M.WHITE, M.LINE)
            M.picture(s, pb, L, TOP, CW, mh, tmp, lossless=True)
            yy = TOP + mh + Inches(0.2)
            aw, ah = Inches(3.6), BOT - yy - Inches(0.4)
            aerial = M.tile_map(ll[0], ll[1], 18, int(aw / 914400 * 200), int(ah / 914400 * 200), "PHOTO2")
            if aerial is not None:
                pa = os.path.join(tmp, "loc_aerial.jpg")
                aerial.save(pa, quality=90)
                M.picture(s, pa, L, yy, aw, ah, tmp)
                M.text(s, L + Inches(0.1), yy + ah - Inches(0.4), Inches(1.6), Inches(0.34), "航照近照", 12, True, M.WHITE,
                       anchor=M.MSO_ANCHOR.MIDDLE)
            rx = L + aw + Inches(0.2)
            rw = R - rx
            rows = [(k, v) for k, v in (("地址", f.get("door")), ("所在商圈", f.get("district")), ("鄰近學校", f.get("school")),
                                        ("面臨道路", f.get("road_w")), ("使用分區", f.get("zone")),
                                        ("座向", f.get("facing"))) if v][:5]
            each = ah // max(1, len(rows))
            for k, v in rows:
                M.box(s, rx, yy, rw, each - Inches(0.1), M.PANEL, M.LINE)
                M.text(s, rx + Inches(0.18), yy + Inches(0.06), rw - Inches(0.3), Inches(0.34), k, 13, True, M.SKY_DK)
                M.text(s, rx + Inches(0.18), yy + Inches(0.4), rw - Inches(0.3), each - Inches(0.55), v,
                       16 if len(v) <= 12 else 13, True, M.INK, anchor=M.MSO_ANCHOR.MIDDLE)
                yy += each
            M.text(s, L, BOT - Inches(0.32), CW, Inches(0.3), "圖資：內政部國土測繪中心（電子地圖、正射影像）", 10, False, M.INK)

    # ================================================================ 地籍圖
    for src in cad_pages:
        M.log("製作：地籍圖")
        s = new("地籍圖")
        png = M.pdf_region_png(src, tmp, M.cadastral_clip, "cadastral")
        ih_ = Inches(6.9)
        M.box(s, L, TOP, CW, ih_, M.WHITE, M.LINE)
        M.picture(s, png, L + Inches(0.08), TOP + Inches(0.08), CW - Inches(0.16), ih_ - Inches(0.16), tmp,
                  mode="contain", lossless=True)
        rows = [(k, v) for k, v in (("地段地號", f.get("lots")), ("使用分區", f.get("zone")), ("建蔽／容積", f.get("bcr_far")),
                                    ("土地坪數", f"{f['land']} 坪" if str(f.get("land", "")).strip() else ""),
                                    ("面寬／深度", f.get("dims_show")), ("比例尺", M.cadastral_scale(src))) if str(v or "").strip()]
        _grid_rows(M, s, rows, L, TOP + ih_ + Inches(0.2), CW, BOT - Inches(0.35))
        M.text(s, L, BOT - Inches(0.3), CW, Inches(0.28), "資料來源：地政事務所地籍圖謄本", 10, False, M.INK)

    # ================================================================ 建物測量成果圖
    for i, src in enumerate(plans):
        M.log("製作：建物測量成果圖")
        s = new("建物測量成果圖" + ("（續）" if i else ""))
        png = M.pdf_region_png(src, tmp, M.plan_clip, "plan") if src.lower().endswith(".pdf") else M.image_plan_png(src, tmp)
        ih_ = Inches(6.6)
        M.box(s, L, TOP, CW, ih_, M.WHITE, M.LINE)
        M.picture(s, png, L + Inches(0.08), TOP + Inches(0.08), CW - Inches(0.16), ih_ - Inches(0.16), tmp,
                  mode="contain", lossless=True)
        rows = [(k, f"{v} 坪") for k, v in (("主建物", f["main"]), ("附屬建物", f["attach"]), ("公共設施", f["public"]))
                if str(v).strip() and not M._zero(v)]
        rows += [(k, f"{v:g} 坪") for k, v in (f.get("floors") or [])[:6]] if len(f.get("floors") or []) > 1 else []
        rows.append(("權狀合計", f"{f['area']} 坪"))
        _grid_rows(M, s, rows, L, TOP + ih_ + Inches(0.2), CW, BOT - Inches(0.35))
        M.text(s, L, BOT - Inches(0.3), CW, Inches(0.28), "資料來源：地政事務所建物測量成果圖", 10, False, M.INK)

    # ================================================================ 格局平面圖
    for i, src in enumerate(layouts[:4]):
        M.log(f"製作：格局平面圖（{os.path.basename(src)}）")
        s = new("格局平面圖" + ("（續）" if i else ""))
        if src.lower().endswith(".pdf"):
            import fitz
            doc = fitz.open(src)
            png = os.path.join(tmp, f"layout_{i}.png")
            doc[0].get_pixmap(matrix=fitz.Matrix(4, 4), alpha=False).save(png)
            doc.close()
        else:
            png = src
        M.box(s, L, TOP, CW, BOT - TOP - Inches(0.45), M.WHITE, M.LINE)
        M.picture(s, png, L + Inches(0.1), TOP + Inches(0.1), CW - Inches(0.2), BOT - TOP - Inches(0.65), tmp,
                  mode="contain", lossless=True)
        M.text(s, L, BOT - Inches(0.35), CW, Inches(0.3), "平面圖為示意，實際以現況及權狀為準", 11, False, M.INK)

    # ================================================================ 照片頁
    M.log("製作：照片頁")
    pages = M.STREET_PHOTO_PAGES if template == "villa" else M.PHOTO_PAGES
    for title, cats, cap in pages:
        items = M.ordered([p for p in chosen if p["cat"] in cats])
        if cover in items and len(items) > 1:
            items = [p for p in items if p is not cover] + [cover]
        for k, chunk in enumerate(_photo_pages(items, CW, BOT - TOP)):
            s = new(title + ("（續）" if k else ""))
            _draw_photos(M, s, chunk, L, TOP, CW, BOT - TOP, tmp)

    # ================================================================ 市場行情
    for p in M.sel_paths(sel, "market", M.market_default(roots, rent)):
        if p.lower().endswith(".pdf"):
            try:
                r = M.crop_104(p, tmp, max_ratio=(Inches(7.2) / CW))
            except Exception:
                r = None
            if not r:
                continue
            img_m, units, src_ = r[0], (None if rent else r[1]), "104 行情通實價網"
        else:
            img_m, units, src_ = p, None, "內政部實價登錄"
        M.log(f"製作：市場行情（{os.path.basename(p)}）")
        s = new("市場行情")
        iw, ih = Image.open(img_m).size
        dh = min(int(CW * ih / iw), Inches(7.4))
        M.box(s, L, TOP, CW, dh + Inches(0.16), M.WHITE, M.LINE)
        M.picture(s, img_m, L + Inches(0.08), TOP + Inches(0.08), CW - Inches(0.16), dh, tmp, mode="contain")
        yy = TOP + dh + Inches(0.35)
        M.gold(s, L, yy, Inches(3.6), Inches(1.0))
        M.text(s, L + Inches(0.2), yy + Inches(0.05), Inches(3.3), Inches(0.35), "本案租金" if rent else "本案開價", 13, True, M.WHITE)
        M.text(s, L + Inches(0.2), yy + Inches(0.38), Inches(3.3), Inches(0.55), f"{f['price']} {f['price_unit']}", 24, True,
               M.WHITE, anchor=M.MSO_ANCHOR.MIDDLE)
        if units:
            rx = L + Inches(3.8)
            M.box(s, rx, yy, R - rx, Inches(1.0), M.WHITE, M.LINE)
            M.text(s, rx + Inches(0.2), yy + Inches(0.05), R - rx - Inches(0.3), Inches(0.35), f"表列 {len(units)} 筆成交單價",
                   12, True, M.SKY_DK)
            M.text(s, rx + Inches(0.2), yy + Inches(0.38), R - rx - Inches(0.3), Inches(0.55),
                   f"{min(units):.1f}～{max(units):.1f} 萬/坪（平均 {sum(units) / len(units):.1f}）", 15, True, M.NUM,
                   anchor=M.MSO_ANCHOR.MIDDLE)
        M.text(s, L, BOT - Inches(0.3), CW, Inches(0.28), f"資料來源：{src_}（提供客戶前請再查證）", 10, False, M.INK)

    # ================================================================ 封底(富住通)
    if back_cover:
        M.log("製作：封底")
        s = new()
        M.box(s, 0, 0, W, H, M.WHITE, radius=False)
        M.gold(s, 0, H - Inches(0.35), W, Inches(0.35), radius=False)
        lv = M.asset(M.ASSET_LOGO_V)
        if lv:
            iw, ih = Image.open(lv).size
            hh = Inches(3.6)
            ww = int(hh * iw / ih)
            s.shapes.add_picture(lv, (W - ww) // 2, Inches(1.8), ww, hh)
        fl_ = M.asset(M.ASSET_FULLLIFE)
        if fl_:
            iw, ih = Image.open(fl_).size
            ww = Inches(3.8)
            s.shapes.add_picture(fl_, (W - ww) // 2, Inches(5.7), ww, int(ww * ih / iw))
        M.text(s, L, Inches(7.0), CW, Inches(0.6), f"{M.COMPANY}　{M.COMPANY_TEL}", 22, True, M.SKY_DK,
               M.PP_ALIGN.CENTER, M.MSO_ANCHOR.MIDDLE)
        ag = [(nm, phones.get(nm, "")) for nm in f["agents"][:3]] if show_agents else []
        for i, (nm, ph_) in enumerate(ag):
            M.text(s, L, Inches(7.9) + Inches(0.55) * i, CW, Inches(0.5), f"{nm}　{ph_}", 18, True, M.INK,
                   M.PP_ALIGN.CENTER, M.MSO_ANCHOR.MIDDLE)
        M.text(s, L, Inches(9.8), CW, Inches(0.5), f"{f['title']}　｜　{f['door']}", 13, True, M.SKY_DK,
               M.PP_ALIGN.CENTER, M.MSO_ANCHOR.MIDDLE)

    M.unique_shape_ids(deck.prs)
    deck.prs.save(out_path)
    return deck.n


def _short_dims(s):
    s = re.sub(r"[（(].*?[）)]", "", str(s or "")).replace(" ", "")
    m = re.search(r"([\d.]+)\D+([\d.]+)", s)
    return f"{m.group(1)}／{m.group(2)} 米" if m else s


def _grid_rows(M, s, rows, x, y, w, bottom):
    """資料兩欄排(標籤＋數值)，高度不超過 bottom"""
    from pptx.util import Inches
    if not rows:
        return
    half = (len(rows) + 1) // 2
    rh = min(Inches(0.62), (bottom - y) // half)
    cw = (w - Inches(0.15)) // 2
    for i, (k, v) in enumerate(rows):
        cx = x + (cw + Inches(0.15)) * (i // half)
        cy = y + rh * (i % half)
        last = k == "權狀合計"
        if last:
            M.gold(s, cx, cy, cw, rh - Inches(0.08))
        else:
            M.box(s, cx, cy, cw, rh - Inches(0.08), M.WHITE, M.LINE)
        M.text(s, cx + Inches(0.15), cy, Inches(1.4), rh - Inches(0.08), k, 12, True, M.WHITE if last else M.SKY_DK,
               anchor=M.MSO_ANCHOR.MIDDLE)
        M.text(s, cx + Inches(1.5), cy, cw - Inches(1.6), rh - Inches(0.08), str(v), 14 if len(str(v)) <= 14 else 11, True,
               M.WHITE if last else M.INK, anchor=M.MSO_ANCHOR.MIDDLE)


def _is_portrait(p):
    from PIL import Image
    try:
        im = Image.open(p["path"])
        w, h = im.size
        try:
            o = im.getexif().get(274)
            if o in (5, 6, 7, 8):
                w, h = h, w
        except Exception:
            pass
        return h > w
    except Exception:
        return False


def _photo_pages(items, w, h):
    """照片分頁：直式兩張一列(列高＝半寬×4/3)、橫式一張一列(列高＝寬×2/3)；依高度一頁塞到放不下為止"""
    from pptx.util import Inches
    gap = Inches(0.16)
    rows, pend = [], []
    for p in items:
        if _is_portrait(p):
            pend.append(p)
            if len(pend) == 2:
                rows.append(pend)
                pend = []
        else:
            rows.append([p])
    if pend:
        rows.append(pend)
    nat = lambda r: int(((w - gap) / 2) * 4 / 3) if len(r) == 2 or _is_portrait(r[0]) else int(w * 2 / 3)
    pages, cur, used = [], [], 0
    for r in rows:
        need = nat(r) + (gap if cur else 0)
        if cur and used + need > h * 1.08:
            pages.append(cur)
            cur, used = [], 0
            need = nat(r)
        cur.append(r)
        used += need
    if cur:
        pages.append(cur)
    return pages


def _draw_photos(M, s, rows, x, y, w, h, tmp):
    from pptx.util import Inches
    gap = Inches(0.16)
    nat = [int(((w - gap) / 2) * 4 / 3) if len(r) == 2 or _is_portrait(r[0]) else int(w * 2 / 3) for r in rows]
    k = (h - gap * (len(rows) - 1)) / sum(nat)          # 依頁高等比例放大/縮小(照片比例不太變形)
    # 放大上限：一般 1.25 倍；只剩一張直式 → 撐滿整頁；只剩一列兩張直式 → 1.6 倍(不要留一大片空白)
    if len(rows) == 1 and len(rows[0]) == 1 and _is_portrait(rows[0][0]):
        k = min(k, (w * 4 / 3) / nat[0], h / nat[0])
    elif len(rows) == 1 and len(rows[0]) == 2:
        k = min(k, 1.6)
    else:
        k = min(k, 1.25)
    tot = sum(int(n_ * k) for n_ in nat) + gap * (len(rows) - 1)
    cy = y + max(0, (h - tot) // 2)                      # 放不滿就上下置中
    for r_i, row in enumerate(rows):
        rh = int(nat[r_i] * k)
        single_portrait = len(row) == 1 and _is_portrait(row[0])
        cw = (w - gap * (len(row) - 1)) // len(row) if not single_portrait else int(rh * 3 / 4)
        x0 = x + (w - cw) // 2 if single_portrait else x
        for c_i, p in enumerate(row):
            cx = x0 + (cw + gap) * c_i
            _, img = M.picture(s, p["path"], cx, cy, cw, rh, tmp)
            cap = p["cap"] if "cap" in p else M.caption_of(p)
            if cap:
                bh = Inches(0.38)
                bw = min(Inches(0.3) + M._text_w(cap, 13), cw - Inches(0.3))
                bx, by = cx + Inches(0.12), cy + rh - bh - Inches(0.12)
                M.box(s, bx, by, bw, bh, M.tone(img, (0.0, 0.75, 0.5, 1.0)), alpha=0.68)
                M.text(s, bx + Inches(0.1), by, bw - Inches(0.15), bh, cap, 13, True, M.WHITE, anchor=M.MSO_ANCHOR.MIDDLE)
        cy += rh + gap
